from pathlib import Path

from docx import Document as DocxDocument

from app.ingestion.models import Document


class DOCXLoader:
    """Load text content and basic metadata from DOCX files."""

    def load(self, file_path: Path) -> list[Document]:
        if file_path.suffix.lower() != ".docx":
            raise ValueError(f"Expected a DOCX file, got: {file_path.name}")

        document = DocxDocument(str(file_path))
        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        properties = document.core_properties

        return [
            Document(
                filename=file_path.name,
                source_type="docx",
                text="\n".join(paragraphs),
                title=properties.title,
                author=properties.author,
                created_at=properties.created,
                modified_at=properties.modified,
            )
        ]
