import csv
from pathlib import Path

from openpyxl import load_workbook

from app.ingestion.models import Document


def _normalize_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _format_row(headers: list[str], values: list[str]) -> str:
    parts: list[str] = []
    for index, header in enumerate(headers):
        value = values[index] if index < len(values) else ""
        if not value:
            continue
        label = header.strip() or f"column_{index + 1}"
        parts.append(f"{label}: {value}")
    return " | ".join(parts)


class TabularLoader:
    """Load traceable row content from CSV and XLSX files."""

    def load(self, file_path: Path) -> list[Document]:
        suffix = file_path.suffix.lower()
        if suffix == ".csv":
            return self._load_csv(file_path)
        if suffix == ".xlsx":
            return self._load_xlsx(file_path)
        raise ValueError(f"Expected a CSV or XLSX file, got: {file_path.name}")

    def _load_csv(self, file_path: Path) -> list[Document]:
        rows: list[Document] = []
        with file_path.open(encoding="utf-8", newline="") as handle:
            reader = csv.reader(handle)
            parsed_rows = list(reader)

        if not parsed_rows:
            return rows

        headers = [_normalize_cell(value) for value in parsed_rows[0]]
        for row_number, raw_row in enumerate(parsed_rows[1:], start=2):
            values = [_normalize_cell(value) for value in raw_row]
            text = _format_row(headers, values)
            if not text:
                continue
            rows.append(
                Document(
                    filename=file_path.name,
                    source_type="csv",
                    text=text,
                    row=row_number,
                )
            )

        return rows

    def _load_xlsx(self, file_path: Path) -> list[Document]:
        rows: list[Document] = []
        workbook = load_workbook(filename=file_path, read_only=True, data_only=True)

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            parsed_rows = [
                [_normalize_cell(value) for value in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            if not parsed_rows:
                continue

            headers = parsed_rows[0]
            for row_number, values in enumerate(parsed_rows[1:], start=2):
                text = _format_row(headers, values)
                if not text:
                    continue
                rows.append(
                    Document(
                        filename=file_path.name,
                        source_type="xlsx",
                        text=text,
                        row=row_number,
                        sheet=sheet_name,
                    )
                )

        workbook.close()
        return rows
