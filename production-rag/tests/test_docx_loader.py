from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.ingestion import Document, DOCXLoader

SAMPLE_DOCX_PATH = Path("data/sample/politica-seguranca-informacao.docx")


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    file_path = tmp_path / "politica-seguranca-informacao.docx"
    document = DocxDocument()
    document.core_properties.title = "Politica de Seguranca da Informacao"
    document.core_properties.author = "Production RAG Corp"
    document.core_properties.created = datetime(2026, 1, 15, 10, 0, tzinfo=UTC)
    document.core_properties.modified = datetime(2026, 2, 1, 12, 30, tzinfo=UTC)
    document.add_paragraph("Use autenticacao multifator em todos os sistemas.")
    document.add_paragraph("Nao compartilhe credenciais por e-mail.")
    document.save(file_path)
    return file_path


def test_docx_loader_extracts_text_and_basic_metadata(sample_docx: Path) -> None:
    documents = DOCXLoader().load(sample_docx)
    content = documents[0]

    assert len(documents) == 1
    assert isinstance(content, Document)
    assert content.source_type == "docx"
    assert content.filename == "politica-seguranca-informacao.docx"
    assert "autenticacao multifator" in content.text
    assert "Nao compartilhe credenciais" in content.text
    assert content.title == "Politica de Seguranca da Informacao"
    assert content.author == "Production RAG Corp"
    assert content.created_at == datetime(2026, 1, 15, 10, 0, tzinfo=UTC)
    assert content.modified_at == datetime(2026, 2, 1, 12, 30, tzinfo=UTC)


def test_docx_loader_rejects_non_docx_files(tmp_path: Path) -> None:
    text_file = tmp_path / "documento.txt"
    text_file.write_text("conteudo", encoding="utf-8")

    with pytest.raises(ValueError, match="Expected a DOCX file"):
        DOCXLoader().load(text_file)


@pytest.mark.skipif(not SAMPLE_DOCX_PATH.exists(), reason="sample DOCX not available")
def test_docx_loader_reads_repository_sample_docx() -> None:
    content = DOCXLoader().load(SAMPLE_DOCX_PATH)[0]

    assert content.filename == SAMPLE_DOCX_PATH.name
    assert content.text
    assert content.title
