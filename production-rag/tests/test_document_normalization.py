from pathlib import Path

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF
from openpyxl import Workbook

from app.ingestion import Document, DOCXLoader, PDFLoader, TabularLoader

SAMPLE_PDF_PATH = Path("data/sample/politica-ferias.pdf")
SAMPLE_DOCX_PATH = Path("data/sample/politica-seguranca-informacao.docx")
SAMPLE_CSV_PATH = Path("data/sample/processo-reembolso-despesas.csv")


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 10, "Pagina 1: politica de ferias corporativa.")
    file_path = tmp_path / "politica-ferias.pdf"
    pdf.output(str(file_path))
    return file_path


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    file_path = tmp_path / "politica-seguranca-informacao.docx"
    document = DocxDocument()
    document.core_properties.title = "Politica de Seguranca da Informacao"
    document.core_properties.author = "Production RAG Corp"
    document.add_paragraph("Use autenticacao multifator em todos os sistemas.")
    document.save(file_path)
    return file_path


@pytest.fixture
def sample_csv(tmp_path: Path) -> Path:
    file_path = tmp_path / "reembolsos.csv"
    file_path.write_text(
        "categoria,descricao,valor\n"
        "transporte,taxi aeroporto,85.50\n",
        encoding="utf-8",
    )
    return file_path


def test_all_loaders_return_normalized_document_type(
    sample_pdf: Path,
    sample_docx: Path,
    sample_csv: Path,
) -> None:
    pdf_documents = PDFLoader().load(sample_pdf)
    docx_documents = DOCXLoader().load(sample_docx)
    csv_documents = TabularLoader().load(sample_csv)

    assert pdf_documents and docx_documents and csv_documents
    assert all(isinstance(document, Document) for document in pdf_documents)
    assert all(isinstance(document, Document) for document in docx_documents)
    assert all(isinstance(document, Document) for document in csv_documents)


def test_pdf_loader_populates_pdf_fields(sample_pdf: Path) -> None:
    document = PDFLoader().load(sample_pdf)[0]

    assert document.source_type == "pdf"
    assert document.filename == "politica-ferias.pdf"
    assert document.page == 1
    assert document.row is None
    assert document.sheet is None
    assert "Pagina 1" in document.text


def test_docx_loader_populates_docx_metadata(sample_docx: Path) -> None:
    document = DOCXLoader().load(sample_docx)[0]

    assert document.source_type == "docx"
    assert document.filename == "politica-seguranca-informacao.docx"
    assert document.page is None
    assert document.row is None
    assert document.title == "Politica de Seguranca da Informacao"
    assert document.author == "Production RAG Corp"
    assert "autenticacao multifator" in document.text


def test_tabular_loader_populates_csv_fields(sample_csv: Path) -> None:
    document = TabularLoader().load(sample_csv)[0]

    assert document.source_type == "csv"
    assert document.filename == "reembolsos.csv"
    assert document.row == 2
    assert document.sheet is None
    assert document.page is None
    assert "transporte" in document.text


def test_xlsx_loader_populates_sheet_and_row(tmp_path: Path) -> None:
    file_path = tmp_path / "sla-suporte.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prioridades"
    sheet.append(["prioridade", "descricao"])
    sheet.append(["P1", "Sistema critico indisponivel"])
    workbook.save(file_path)

    document = TabularLoader().load(file_path)[0]

    assert document.source_type == "xlsx"
    assert document.sheet == "Prioridades"
    assert document.row == 2
    assert "P1" in document.text


@pytest.mark.skipif(not SAMPLE_PDF_PATH.exists(), reason="sample PDF not available")
@pytest.mark.skipif(not SAMPLE_DOCX_PATH.exists(), reason="sample DOCX not available")
@pytest.mark.skipif(not SAMPLE_CSV_PATH.exists(), reason="sample CSV not available")
def test_repository_samples_use_normalized_document_schema() -> None:
    pdf_documents = PDFLoader().load(SAMPLE_PDF_PATH)
    docx_documents = DOCXLoader().load(SAMPLE_DOCX_PATH)
    csv_documents = TabularLoader().load(SAMPLE_CSV_PATH)

    for document in [*pdf_documents, *docx_documents, *csv_documents]:
        assert isinstance(document, Document)
        assert document.filename
        assert document.source_type in {"pdf", "docx", "csv", "xlsx"}
        assert document.text
